#!/usr/bin/env python3
"""
analisar_job.py — Serviço headless de análise de documentos
════════════════════════════════════════════════════════════
Uso:   python3 analisar_job.py <pasta_raiz>
Saída: JSON no stdout  |  Logs de progresso no stderr

Adaptado de analisador.py (Studio Fiscal v2.0)
⚠  SOMENTE LEITURA — nunca altera, move ou exclui arquivos.
"""

import os
import re
import sys
import json
import unicodedata
from pathlib import Path
from datetime import datetime

# ── Saída sempre em UTF-8 ─────────────────────────────────────────────────────
# No Windows, stdout/stderr redirecionados (pipe) usam a codepage ANSI (cp1252)
# e quebram em caracteres como "≥" ("'charmap' codec can't encode character").
# O .exe do PyInstaller IGNORA PYTHONUTF8/PYTHONIOENCODING, então não dá pra
# depender de variável de ambiente: reconfigura aqui, antes de qualquer print.
for _fluxo in (sys.stdout, sys.stderr):
    try:
        if _fluxo is not None:
            _fluxo.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# ── Importações com tratamento de erro claro ──────────────────────────────────
try:
    import pdfplumber
except ImportError:
    print(json.dumps({"erro": "pdfplumber não instalado. Execute: pip install pdfplumber"}))
    sys.exit(1)


# ══════════════════════════════════════════════════════════════════════════════
#  LIMIARES DE NEGÓCIO (ajustar aqui se as regras mudarem)
# ══════════════════════════════════════════════════════════════════════════════
LIMIAR_FASTPASS = 200_000_000
LIMIAR_SUPPLY   = 30_000_000
LIMIAR_FINAP    = 50_000_000          # + CNAE principal precisa constar na lista Finap

# Regularize: olha SÓ o documento de capacidade de pagamento (PGFN). Se a dívida
# total (PGFN + RFB) estiver ACIMA deste valor, está DENTRO; senão, FORA.
# A classificação (A/B/C/D/omisso) NÃO entra mais na regra (só é exibida).
REGULARIZE_LIMIAR = 600_000

PASTA_CNPJ = ["B - Documentos", "5 - Cliente", "16 - CNPJ"]

# ══════════════════════════════════════════════════════════════════════════════
#  ELEMENTOS DE ANÁLISE
#  "tipo" decide qual processador roda (ver PROCESSADORES mais abaixo)
# ══════════════════════════════════════════════════════════════════════════════
ELEMENTOS: list[dict] = [
    # ── CONTRATO ─────────────────────────────────────────────────────────────
    {
        "id": "contrato_existe",
        "categoria": "Contrato",
        "campo": "Contrato Assinado",
        "descricao": "PDF de contrato em A - Contrato › 1 - Contrato Studio",
        "pasta": ["A - Contrato", "1 - Contrato Studio"],
        "extensao": [".pdf"],
        "tipo": "existencia",
    },
    {
        "id": "procuracao_existe",
        "categoria": "Contrato",
        "campo": "Procuração",
        "descricao": "Definido exclusivamente pelo Project (não depende do arquivo)",
        "tipo": "manual",
    },
    {
        "id": "honorarios",
        "categoria": "Contrato",
        "campo": "Honorários (%)",
        "descricao": "Percentual de honorários — sempre da cláusula 5.1 do contrato",
        "pasta": ["A - Contrato"],
        "arquivo_contem": ["Contrato"],
        "extensao": [".pdf"],
        "tipo": "honorarios",
    },
    {
        "id": "gross_up",
        "categoria": "Contrato",
        "campo": "Gross Up",
        "descricao": "Contrato prevê gross up? (sempre 14,25%)",
        "pasta": ["A - Contrato"],
        "arquivo_contem": ["Contrato"],
        "extensao": [".pdf"],
        "tipo": "gross_up",
    },
    {
        "id": "passivo",
        "categoria": "Contrato",
        "campo": "Passivo",
        "descricao": "Menção a passivos a serem realizados",
        "pasta": ["A - Contrato"],
        "arquivo_contem": ["Contrato"],
        "extensao": [".pdf"],
        "regex": [r"(passivo[s]?[^.\n]{0,300})"],
        "tipo": "booleano",
    },

    # ── DATA DE ABERTURA (sempre da Matriz) ──────────────────────────────────
    {
        "id": "data_abertura",
        "categoria": "Data de Abertura",
        "campo": "Data de Abertura",
        "descricao": "Data de constituição do CNPJ Matriz",
        "pasta": PASTA_CNPJ,
        "arquivo_contem": ["CNPJ"],
        "arquivo_nao_contem": ["QSA", "Optantes", "Consulta"],
        "extensao": [".pdf"],
        "assinatura": r"COMPROVANTE\s+DE\s+INSCRI",
        "tipo": "data_abertura",
    },

    # ── SIMPLES NACIONAL ─────────────────────────────────────────────────────
    {
        "id": "simples_nacional",
        "categoria": "Simples Nacional",
        "campo": "Simples Nacional",
        "descricao": "Situação no Simples Nacional e no SIMEI",
        "pasta": PASTA_CNPJ,
        "arquivo_contem_qualquer": ["Optantes", "Simples"],
        "extensao": [".pdf"],
        "assinatura": r"Simples\s+Nacional",
        "tipo": "simples",
    },

    # ── QSA ──────────────────────────────────────────────────────────────────
    {
        "id": "socios",
        "categoria": "QSA",
        "campo": "Sócios (QSA)",
        "descricao": "Relatório fiscal: nomes dos sócios, CPF/CNPJ e evidências",
        "pasta": ["B - Documentos", "2 - Fiscal"],
        "arquivo_contem_qualquer": ["SituacaoFiscal", "SituaçãoFiscal", "Situacao Fiscal", "Situação Fiscal"],
        "extensao": [".pdf"],
        "assinatura": r"(?:s[oó]cios|quadro\s+societ[aá]rio)",
        "tipo": "qsa",
    },

    # ── REGULARIZE ───────────────────────────────────────────────────────────
    {
        "id": "regularize",
        "categoria": "Regularize",
        "campo": "Regularize",
        "descricao": "Perfil Regularize (capacidade de pagamento PGFN) ou sem cadastro",
        "pasta": ["B - Documentos", "2 - Fiscal"],
        "arquivo_contem_qualquer": ["PGFN", "Procuradoria", "Capag", "Capacidade"],
        "extensao": [".pdf"],
        "assinatura": r"capacidade\s+de\s+pagamento",
        "tipo": "regularize",
    },

    # ── PERFIL PREVIDENCIÁRIO ────────────────────────────────────────────────
    {
        "id": "perfil",
        "categoria": "Perfil Prev.",
        "campo": "Perfil Previdenciário",
        "descricao": "Qualquer ano acima do limiar = É PERFIL. Lista todos os anos.",
        "pasta": ["B - Documentos", "6 - Cruzamentos"],
        "arquivo_contem": ["perfil"],
        "extensao": [".pdf", ".xlsx", ".xls", ".txt"],
        "tipo": "perfil",
    },

    # ── FATURAMENTO ──────────────────────────────────────────────────────────
    {
        "id": "faturamento",
        "categoria": "Faturamento",
        "campo": "Faturamento por Ano",
        "descricao": "Total por ano (arquivos ECF) + média dos anos",
        "pasta": ["B - Documentos", "1 - Contábil", "5 - ECF"],

        "extensao": [".txt", ".pdf", ".xlsx"],
        "tipo": "faturamento",
    },
    {
        "id": "fastpass",
        "categoria": "Faturamento",
        "campo": "Perfil Fastpass (> R$ 200M)",
        "descricao": "Faturamento médio acima de R$ 200 milhões?",
        "tipo": "derivado",
        "limiar": LIMIAR_FASTPASS,
    },
    {
        "id": "supply",
        "categoria": "Faturamento",
        "campo": "Supply (≥ R$ 30M)",
        "descricao": "Faturamento médio a partir de R$ 30 milhões (abaixo disso desqualifica)",
        "tipo": "derivado",
        "limiar": LIMIAR_SUPPLY,
    },
    {
        "id": "finap",
        "categoria": "Faturamento",
        "campo": "Perfil Finap - Lei do Bem",
        "descricao": "Faturamento médio acima de R$ 50 milhões e CNAE principal na lista da planilha",
        "tipo": "derivado",
        "limiar": LIMIAR_FINAP,
        "exige_industria": True,
    },
]


# ══════════════════════════════════════════════════════════════════════════════
#  BUSCA DE ARQUIVOS
# ══════════════════════════════════════════════════════════════════════════════
def _normalizar(texto):
    return ''.join(c for c in unicodedata.normalize('NFKD', texto or '') if not unicodedata.combining(c)).casefold()

def _contem_todos(nome: str, termos: list[str]) -> bool:
    n = _normalizar(nome)
    return all(_normalizar(t) in n for t in termos)


def _contem_qualquer(nome: str, termos: list[str]) -> bool:
    n = _normalizar(nome)
    return any(_normalizar(t) in n for t in termos)


def _nao_contem(nome: str, termos: list[str]) -> bool:
    n = _normalizar(nome)
    return not any(_normalizar(t) in n for t in termos)


def _varrer_pasta(raiz: str, pastas_alvo: list[str], extensoes: set[str]):
    """Gera todos os arquivos cujo caminho passa pelas pastas-alvo."""
    for dirpath, dirs, filenames in os.walk(raiz):
        dirs.sort()
        partes = [_normalizar(p) for p in Path(dirpath).parts]
        # Ordered full components: 2 - Fiscal must not match 12 - Fiscal.
        pos = 0
        for parte in partes:
            if pos < len(pastas_alvo) and parte == _normalizar(pastas_alvo[pos]):
                pos += 1
        if pos != len(pastas_alvo):
            continue
        for fname in sorted(filenames):
            if extensoes and Path(fname).suffix.lower() not in extensoes:
                continue
            yield os.path.join(dirpath, fname)


def encontrar_arquivos(raiz: str, elemento: dict) -> list[str]:
    """Busca por NOME de arquivo."""
    pastas_alvo    = elemento.get("pasta", [])
    arq_contem     = elemento.get("arquivo_contem", [])
    arq_qualquer   = elemento.get("arquivo_contem_qualquer", [])
    arq_nao_contem = elemento.get("arquivo_nao_contem", [])
    extensoes      = {e.lower() for e in elemento.get("extensao", [])}
    resultados: list[str] = []

    _log(f"      [busca] pasta={pastas_alvo} | contem={arq_contem or arq_qualquer} | ext={list(extensoes) or 'qualquer'}")

    for caminho in _varrer_pasta(raiz, pastas_alvo, extensoes):
        fname = os.path.basename(caminho)
        if arq_contem and not _contem_todos(fname, arq_contem):
            continue
        if arq_qualquer and not _contem_qualquer(fname, arq_qualquer):
            continue
        if arq_nao_contem and not _nao_contem(fname, arq_nao_contem):
            continue
        _log(f"      [✓] {caminho}")
        resultados.append(caminho)

    if not resultados:
        _log("      [✗] nenhum arquivo encontrado pelo nome")
    return resultados


def candidatos(raiz: str, elemento: dict):
    """
    Gera arquivos candidatos: primeiro os que casam pelo NOME; depois (preguiçoso)
    os demais arquivos da pasta cujo CONTEÚDO bate com a assinatura do elemento.
    Assim funciona mesmo quando o arquivo tem nome genérico (ex.: 'pdf2.pdf').
    """
    vistos: set[str] = set()
    for a in encontrar_arquivos(raiz, elemento):
        vistos.add(a)
        yield a

    assinatura = elemento.get("assinatura")
    if not assinatura:
        return
    extensoes = {e.lower() for e in elemento.get("extensao", [".pdf"])}
    for a in _varrer_pasta(raiz, elemento.get("pasta", []), extensoes):
        if a in vistos:
            continue
        if re.search(assinatura, extrair_texto(a), re.IGNORECASE):
            _log(f"      [conteúdo] {a}")
            vistos.add(a)
            yield a

# ══════════════════════════════════════════════════════════════════════════════
#  OCR VIA TESSERACT
# ══════════════════════════════════════════════════════════════════════════════
def _garantir_tesseract() -> bool:
    """Localiza o Tesseract: portátil (empacotado), variável de ambiente,
    instalação padrão do Windows, ou no PATH (Linux/Docker)."""
    import shutil

    def _usar(caminho: str) -> bool:
        if not os.path.exists(caminho):
            return False
        try:
            import pytesseract
            pytesseract.pytesseract.tesseract_cmd = caminho
        except ImportError:
            pass
        return True

    # 1) Variável de ambiente explícita (setada pelo launcher ao lado do .exe portátil)
    env_cmd = os.environ.get("TESSERACT_CMD")
    if env_cmd and _usar(env_cmd):
        return True

    # 2) Tesseract portátil ao lado do próprio executável (empacotado com PyInstaller)
    #    sys.executable = caminho do analisar_job.exe quando congelado
    if getattr(sys, "frozen", False):
        base = os.path.dirname(sys.executable)
        for rel in (r"tesseract\tesseract.exe", r"..\tesseract\tesseract.exe"):
            if _usar(os.path.join(base, rel)):
                return True

    # 3) Docker / Linux: tesseract instalado via apt (no PATH)
    if shutil.which("tesseract"):
        return True

    # 4) Windows: instalação padrão do instalador oficial
    if os.name == 'nt':
        for p in [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ]:
            if _usar(p):
                return True

    _log("      [ocr] Tesseract não encontrado. Instale via:")
    _log("            Linux/Docker: apt-get install tesseract-ocr tesseract-ocr-por")
    _log("            Windows: winget install UB-Mannheim.TesseractOCR")
    return False


def _ocr_pdf(caminho: str, paginas=None) -> str:
    _log("      [ocr] PDF sem texto — ativando OCR Tesseract...")
    if not _garantir_tesseract():
        return ""
    try:
        import pymupdf as fitz  # nova API (evita warning no stdout)
    except ImportError:
        try:
            import fitz  # fallback versão antiga
        except ImportError:
            _log("      [ocr] pymupdf não instalado")
            return ""
    try:
        from PIL import Image
    except ImportError:
        _log("      [ocr] Pillow não instalado")
        return ""
    try:
        import pytesseract
    except ImportError:
        pytesseract = None
    import io

    try:
        doc = fitz.open(caminho)
        partes: list[str] = []
        for i, page in enumerate(doc):
            if paginas is not None and i not in paginas:
                continue
            pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5))
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            if pytesseract is None:
                import tempfile, subprocess, shutil
                with tempfile.TemporaryDirectory(prefix='analise-ocr-') as tmp:
                    imagem = os.path.join(tmp, 'pagina.png')
                    img.save(imagem)
                    cmd = os.environ.get('TESSERACT_CMD') or shutil.which('tesseract')
                    if not cmd: return ''
                    proc = subprocess.run([cmd, imagem, 'stdout', '-l', 'por'], capture_output=True, text=True, encoding='utf-8', timeout=120)
                    if proc.returncode:
                        proc = subprocess.run([cmd, imagem, 'stdout', '-l', 'eng'], capture_output=True, text=True, encoding='utf-8', timeout=120)
                    txt = proc.stdout if proc.returncode == 0 else ''
            else:
                try:
                    txt = pytesseract.image_to_string(img, lang="por", timeout=120)
                except Exception:
                    txt = pytesseract.image_to_string(img, timeout=120)
            partes.append(txt)
            _log(f"      [ocr] página {i+1}: {len(txt)} chars")
        doc.close()
        resultado = "\n".join(partes)
        if resultado.strip():
            _log(f"      [ocr] total {len(resultado)} chars extraídos")
        return resultado
    except Exception as e:
        _log(f"      [ocr] ERRO: {e}")
        return ""


# ══════════════════════════════════════════════════════════════════════════════
#  EXTRAÇÃO DE TEXTO
# ══════════════════════════════════════════════════════════════════════════════
def _extrair_texto_raw(caminho: str) -> str:
    ext  = Path(caminho).suffix.lower()
    nome = os.path.basename(caminho)
    _log(f"      [lendo] {nome}")
    try:
        if ext == ".pdf":
            partes = []
            with pdfplumber.open(caminho) as pdf:
                for i, page in enumerate(pdf.pages):
                    try:
                        texto = page.extract_text() or ''
                    except Exception:
                        texto = ''
                    if not texto.strip():
                        try:
                            import pymupdf
                            with pymupdf.open(caminho) as doc:
                                texto = doc[i].get_text()
                        except Exception:
                            pass
                    if not texto.strip():
                        texto = _ocr_pdf(caminho, {i})
                    partes.append(f"\n[Página {i+1}]\n{texto}")
            return "\n".join(partes) if any(re.search(r"[a-zA-Z]", x.split(']')[-1]) for x in partes) else ''

        elif ext == ".txt":
            for enc in ("utf-8", "latin-1", "cp1252"):
                try:
                    with open(caminho, "r", encoding=enc, errors="strict") as f:
                        texto = f.read()
                    _log(f"      [txt] {len(texto)} chars")
                    return texto
                except UnicodeDecodeError:
                    continue
            with open(caminho, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        elif ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
                linhas: list[str] = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        linha = " ".join(str(c) for c in row if c is not None)
                        if linha.strip():
                            linhas.append(linha)
                wb.close()
                texto = "\n".join(linhas)
                _log(f"      [xlsx] {len(linhas)} linhas")
                return texto
            except Exception as e:
                _log(f"      [xlsx] ERRO: {e}")
                return ""

        else:
            for enc in ("utf-8", "latin-1", "cp1252"):
                try:
                    with open(caminho, "r", encoding=enc, errors="strict") as f:
                        return f.read()
                except (UnicodeDecodeError, OSError):
                    continue
            return ""

    except Exception as e:
        _log(f"      [ERRO] {e}")
        return ""



_CACHE_TEXTO: dict[tuple, str] = {}


def extrair_texto(caminho: str) -> str:
    """Extrai texto com cache (o mesmo arquivo pode ser lido por vários elementos)."""
    st = os.stat(caminho)
    chave = (caminho, st.st_mtime_ns, st.st_size)
    if chave not in _CACHE_TEXTO:
        _CACHE_TEXTO[chave] = _extrair_texto_raw(caminho)
    return _CACHE_TEXTO[chave]


# ══════════════════════════════════════════════════════════════════════════════
#  HELPERS NUMÉRICOS (formato brasileiro)
# ══════════════════════════════════════════════════════════════════════════════
_NUM = r"(?:\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+(?:[.,]\d+)?)"


def _brl(s: str):
    """'1.234.567,89' -> 1234567.89 | '45,3' -> 45.3 | '45.3' -> 45.3 | None se inválido."""
    s = (s or "").replace("R$", "").strip()
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", s):
        return float(s.replace(".", "").replace(",", "."))
    if re.fullmatch(r"\d+,\d+", s):
        return float(s.replace(",", "."))
    if re.fullmatch(r"\d+(?:\.\d+)?", s):
        return float(s)
    return None


def _fmt_brl(v: float) -> str:
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_num(v: float) -> str:
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _trecho(txt: str, ini: int, fim: int) -> str:
    return "..." + txt[max(0, ini - 60):min(len(txt), fim + 160)].replace("\n", " ").strip() + "..."


def _r(valor="", trecho="", arquivos=None, status=None, dados=None) -> dict:
    return {
        "valor": valor, "trecho": trecho, "arquivos": arquivos or [],
        "status": status or ("ok" if valor else "nao_encontrado"),
        "dados": dados or {},
    }


# ══════════════════════════════════════════════════════════════════════════════
#  PROCESSADORES
#  Assinatura: (raiz, elemento, ctx) -> dict (ver _r)
#  ctx: dicionário compartilhado entre elementos (ex.: CNAE principal)
# ══════════════════════════════════════════════════════════════════════════════

# ── Contrato: existência / manual ────────────────────────────────────────────
def _proc_existencia(raiz, elem, ctx):
    arquivos = encontrar_arquivos(raiz, elem)
    if arquivos:
        return _r(f"Sim — {os.path.basename(arquivos[0])}", arquivos=arquivos)
    return _r()


def _proc_manual(raiz, elem, ctx):
    # Procuração: vem exclusivamente do Project. Fica em branco por enquanto.
    return _r(status="em_branco")


# ── Honorários: SEMPRE da cláusula 5.1 ───────────────────────────────────────
_RE_INI_51_LINHA = re.compile(r"(?mi)^[ \t]*(?:cl[aá]usula[ \t]*)?5\.1(?!\d)")
_RE_INI_51_QQ    = re.compile(r"(?<![\d.,])5\.1(?!\d)")
_RE_FIM_51       = re.compile(
    r"(?mi)^[ \t]*(?:(?:cl[aá]usula[ \t]*)?5\.2(?!\d)"
    r"|cl[aá]usula[ \t]*(?:6|sexta)\b|6[.\-–)](?:\d|\s))"
)
_RE_PERCENT = re.compile(r"(\d{1,3}(?:[.,]\d{1,4})?)\s*%")
_RE_POR_EXTENSO = re.compile(r"\(\s*([a-zçãõáéíóúâêô ]{3,60}?)\s+por\s+cento\s*\)", re.IGNORECASE)


def _blocos_clausula_51(texto: str):
    """Gera candidatos de bloco da cláusula 5.1 (primeiro os que começam em linha própria)."""
    inicios = [m.start() for m in _RE_INI_51_LINHA.finditer(texto)]
    inicios += [m.start() for m in _RE_INI_51_QQ.finditer(texto) if m.start() not in inicios]
    for ini in inicios:
        fim = ini + 1800
        m_fim = _RE_FIM_51.search(texto, ini + 4)
        if m_fim and m_fim.start() < fim:
            fim = m_fim.start()
        yield ini, texto[ini:fim]


def _percentual_do_bloco(bloco: str):
    for m in _RE_PERCENT.finditer(bloco):
        num = m.group(1)
        antes = bloco[max(0, m.start() - 60):m.start()]
        if re.search(r"gross[\s\-]?up[^%]{0,40}$", antes, re.IGNORECASE):
            continue                       # percentual do gross up, não dos honorários
        if _brl(num) == 14.25:
            continue
        return num + "%", m
    m = _RE_POR_EXTENSO.search(bloco)
    if m:
        return m.group(1).strip() + " por cento", m
    return None, None


def _proc_honorarios(raiz, elem, ctx):
    arquivos = list(candidatos(raiz, elem))
    for arq in arquivos:
        texto = extrair_texto(arq)
        for ini, bloco in _blocos_clausula_51(texto):
            val, m = _percentual_do_bloco(bloco)
            if val:
                _log(f"      → cláusula 5.1: {val}")
                return _r(val, _trecho(texto, ini + m.start(), ini + m.end()), [arq])

    # Fallback (marcado para conferência): primeiro % próximo de 'honorários'
    for arq in arquivos:
        texto = extrair_texto(arq)
        m = re.search(r"(?:honor[aá]rios|remunera[çc][aã]o)[\s\S]{0,500}?(\d{1,3}(?:[.,]\d{1,4})?)\s*%", texto, re.IGNORECASE)
        if m:
            return _r(f"{m.group(1)}% (cláusula 5.1 não localizada — conferir)",
                      _trecho(texto, m.start(), m.end()), [arq], status="ok")
    return _r(arquivos=arquivos[:1], status="sem_dado" if arquivos else "nao_encontrado")


# ── Gross up: só Sim / Não (sempre 14,25%) ───────────────────────────────────
def _proc_gross_up(raiz, elem, ctx):
    arquivos = list(candidatos(raiz, elem))
    if not arquivos:
        return _r()
    legiveis = []
    for arq in arquivos:
        texto = extrair_texto(arq)
        if texto.strip(): legiveis.append(arq)
        m = re.search(r"gross[\s\-]?up|14[,\.]25\s*%", texto, re.IGNORECASE)
        if m and not re.search(r"(?:n[aã]o|sem)\s+(?:se\s+aplica\s+)?gross[\s-]?up", texto, re.I):
            return _r("Sim (14,25%)", _trecho(texto, m.start(), m.end()), [arq])
    if not legiveis:
        return _r("Inconclusivo — contrato sem texto legível", arquivos=arquivos, status="inconclusivo")
    return _r("Não", arquivos=legiveis, status="ok")


# ── Booleano genérico (Passivo) ──────────────────────────────────────────────
def _proc_booleano(raiz, elem, ctx):
    arquivos = list(candidatos(raiz, elem))
    for arq in arquivos:
        texto = extrair_texto(arq)
        for p in elem.get("regex", []):
            m = re.search(p, texto, re.IGNORECASE)
            if m and not re.search(r"(?:sem|n[aã]o\s+(?:h[aá]|existem?|possui))\s+passivos?", texto[max(0,m.start()-50):m.end()], re.I):
                return _r("Sim (mencionado no contrato)", _trecho(texto, m.start(), m.end()), [arq])
    return _r(arquivos=arquivos[:1], status="sem_dado" if arquivos else "nao_encontrado")


# ── Data de abertura: bloco da MATRIZ ────────────────────────────────────────
def _bloco_matriz(texto: str) -> str:
    partes = re.split(r"(?=N[ÚU]MERO\s+DE\s+INSCRI\S{1,3}O)", texto, flags=re.I)
    for bloco in partes:
        if re.search(r"\bMATR[IE]Z\b", bloco, re.I) and not re.search(r"\bFILIAL\b", bloco[:500], re.I):
            return bloco
    # No fallback to arbitrary text or to /0001: explicit evidence of matriz is mandatory.
    return ''


def _proc_data_abertura(raiz, elem, ctx):
    observacoes, arquivos = [], []
    candidatos_ordenados = sorted(set(candidatos(raiz, elem)), key=lambda x: ('matriz' not in _normalizar(os.path.basename(x)), x))
    for arq in candidatos_ordenados:
        texto = extrair_texto(arq)
        bloco = _bloco_matriz(texto)
        if not bloco:
            continue
        arquivos.append(arq)
        m = re.search(r"DATA\s+DE\s+(?:CONSTITUI\S{1,3}O|ABERTURA)[\s\S]{0,200}?(\d{2}[/.-]\d{2}[/.-]\d{4})", bloco, re.I)
        data = None
        if m:
            try:
                dt = datetime.strptime(re.sub(r'[.-]', '/', m.group(1)), '%d/%m/%Y')
                if dt <= datetime.now(): data = dt.strftime('%d/%m/%Y')
            except ValueError:
                pass
        mc = re.search(r"ATIVIDADE\s+ECON[ÔO]MICA\s+PRINCIPAL\s*\n?\s*(\d{2}\.?\d{2}-?\d[/-]?\d{2})\s*[-–]\s*([^\n]+)", bloco, re.I)
        cnpj = re.search(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", bloco)
        observacoes.append({"data":data,"cnae":mc.group(1) if mc else None,"cnpj":cnpj.group(0) if cnpj else None,"arquivo":arq,"trecho":bloco[:1800]})
    if not observacoes or not any(x['data'] and x['cnae'] and x['cnpj'] for x in observacoes):
        for arq in _varrer_pasta(raiz, ['B - Documentos','2 - Fiscal'], {'.pdf','.txt'}):
            texto = extrair_texto(arq)
            bloco = re.search(r"Dados\s+Cadastrais\s+da\s+Matriz([\s\S]*?)(?=S[oó]cios\s+e\s+Administradores|Certid[aã]o|Diagn[oó]stico|\Z)", texto, re.I)
            if not bloco: continue
            t = bloco.group(1)
            md = re.search(r"Data\s+de\s+Abertura\s*:\s*(\d{2}/\d{2}/\d{4})",t,re.I)
            mc = re.search(r"CNAE\s*:\s*(\d{4}-\d/\d{2})",t,re.I)
            mj = re.search(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}",t)
            if md:
                try:
                    if datetime.strptime(md.group(1),'%d/%m/%Y') > datetime.now(): continue
                except ValueError: continue
            if md or mc:
                arquivos.append(arq)
                observacoes.append({'data':md.group(1) if md else None,'cnae':mc.group(1) if mc else None,'cnpj':mj.group(0) if mj else None,'arquivo':arq,'trecho':bloco.group(0)})
    datas = {x['data'] for x in observacoes if x['data']}
    cnaes = {re.sub(r'\D','',x['cnae']) for x in observacoes if x['cnae']}
    cnpjs = {x['cnpj'] for x in observacoes if x['cnpj']}
    conflito = len(datas)>1 or len(cnaes)>1 or len(cnpjs)>1
    if conflito or not observacoes:
        return _r('Inconclusivo — matriz ausente ou documentos de matriz conflitantes', arquivos=arquivos,
                  status='inconclusivo', dados={"evidencias":observacoes})
    fonte = next((x for x in observacoes if x['cnae']), observacoes[0])
    ctx.update(cnae_principal=fonte['cnae'], cnae_arquivo=fonte['arquivo'], cnae_trecho=fonte['trecho'])
    if len(cnpjs)==1: ctx['cnpj_matriz'] = next(iter(cnpjs))
    return _r(next(iter(datas)) if datas else 'Data não identificada', arquivos=arquivos,
              dados={"cnae_principal":fonte['cnae'], "cnpj_matriz":ctx.get('cnpj_matriz'), "evidencias":observacoes})


# ── Simples Nacional ─────────────────────────────────────────────────────────
def _proc_simples(raiz, elem, ctx):
    situacoes, periodos, evidencias, arquivos = [], [], [], []
    sem_historico = []
    hoje = datetime.now()
    for arq in candidatos(raiz, elem):
        texto = extrair_texto(arq)
        reconhecido = False
        for label in ("Simples Nacional", "SIMEI"):
            m = re.search(r"Situa\S{1,3}o\s+no\s+" + label + r"\s*:\s*([^\n\r]+)", texto, re.I)
            if m:
                situacoes.append(label + ": " + m.group(1).strip())
                evidencias.append({"arquivo": arq, "trecho": m.group(0)})
                reconhecido = True
        neg = re.search(r"(?:op\S{1,4}es|op[çc][oõ]es)\s+pelo\s+Simples\s+Nacional\s+em\s+Per\S{1,3}odos\s+Anteriores\s*:\s*N[aã]o\s+Existem", texto, re.I)
        if neg:
            sem_historico.append(arq)
            evidencias.append({'arquivo':arq,'trecho':neg.group(0)})
            reconhecido = True
        # Work line by line, keeping Simples and SIMEI tables separate.
        secao = None
        for linha in texto.splitlines():
            if re.search(r"SIMEI", linha, re.I):
                secao = "simei"
            elif re.search(r"Simples Nacional", linha, re.I):
                secao = "simples"
            if secao != "simples":
                continue
            if re.search(r"agendament|eventos? futuros?", linha, re.I):
                secao = None
            if secao != 'simples': continue
            datas = re.findall(r"\d{2}/\d{2}/\d{4}", linha)
            if len(datas) < 2 and not (datas and re.search(r"optante.*desde|data\s+(?:de\s+)?op[çc][aã]o|atual|vigente", linha, re.I)):
                continue
            try:
                inicio = datetime.strptime(datas[0], '%d/%m/%Y')
                fim = datetime.strptime(datas[1], '%d/%m/%Y') if len(datas)>1 else hoje
            except ValueError:
                continue
            if inicio > fim or inicio > hoje:
                continue
            anos = list(range(inicio.year, min(fim.year, hoje.year)+1))
            periodos.append({"inicio": datas[0], "fim": datas[1] if len(datas)>1 else "atual", "anos": anos, "arquivo": arq, "trecho": linha})
            evidencias.append({"arquivo": arq, "trecho": linha})
            reconhecido = True
        if reconhecido:
            arquivos.append(arq)
    anos = sorted({ano for periodo in periodos for ano in periodo['anos']})
    linhas = list(dict.fromkeys(situacoes))
    if sem_historico and not anos:
        linhas.append('Sem períodos anteriores de opção pelo Simples Nacional, conforme consulta.')
    linhas.append("Já foi optante. Anos documentados: " + ", ".join(map(str, anos)) if anos else
                  "Nenhum ano anterior documentado como optante." if sem_historico else
                  "Histórico de anos no Simples não identificado; não inferido da situação atual.")
    return _r("\n".join(linhas) if arquivos else "", arquivos=arquivos,
              dados={"anos_optante": anos, "periodos": periodos, "sem_periodos_anteriores": bool(sem_historico), "evidencias": evidencias})


# ── QSA ──────────────────────────────────────────────────────────────────────
_RE_SOCIO = re.compile(
    r"Nome(?:\s*/\s*Nome\s*Empresarial)?\s*:\s*(?P<n>[^\n\r]+)"
    r"(?:(?!Nome(?:\s*/\s*Nome\s*Empresarial)?\s*:)[\s\S]){0,250}?"
    r"Qualifica\S{1,3}o\s*:\s*(?P<q>[^\n\r]+)",
    re.IGNORECASE)


def _proc_qsa(raiz, elem, ctx):
    socios, evidencias, arquivos = [], [], []
    vistos = set()
    for arq in _varrer_pasta(raiz, elem["pasta"], {".pdf", ".txt", ".xlsx"}):
        texto = extrair_texto(arq)
        # Restrict to the corporate-partner section, never arbitrary CPF/CNPJ in a report.
        sec = re.search(r"(?:quadro\s+(?:de\s+)?s[oó]cios[^\n]*|quadro\s+societ[aá]rio|s[oó]cios\s+e\s+administradores)([\s\S]*?)(?=\n\s*(?:Certid[aã]o|Diagn[oó]stico|Pend[eê]ncias|D[eé]bitos|Situa[çc][aã]o Fiscal)\b|\Z)", texto, re.I)
        if not sec:
            continue
        bloco = sec.group(1)
        for linha in bloco.splitlines():
            m = re.search(r"(?P<doc>\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}|\d{3}\.\d{3}\.\d{3}-\d{2}|(?<!\d)(?:\d{14}|\d{11})(?!\d))", linha)
            if not m:
                continue
            documento = m.group('doc')
            nome = (linha[:m.start()] + ' ' + linha[m.end():]).strip(' :|-')
            if re.match(r"^(?:CNPJ|CPF|Respons[aá]vel)\s*:", linha, re.I):
                continue
            qualificacao = re.search(r"\b(PRESIDENTE|DIRETOR(?:A)?|ADMINISTRADOR(?:A)?|SOCIO(?:-ADMINISTRADOR)?|S[ÓO]CIO(?:-ADMINISTRADOR)?|CONSELHEIRO|TITULAR)\b", nome, re.I)
            cargo = nome[qualificacao.start():] if qualificacao else ''
            if qualificacao: nome = nome[:qualificacao.start()].strip()
            if not nome or (nome, documento) in vistos:
                continue
            vistos.add((nome, documento))
            socios.append({"nome": nome, "qualificacao": cargo, "documento": documento, "tipo": "CNPJ" if len(re.sub(r"\D", "", documento)) == 14 else "CPF"})
            evidencias.append({"arquivo": arq, "trecho": linha})
            if arq not in arquivos:
                arquivos.append(arq)
    return _r("\n".join(f"{x['nome']} | {x['tipo']}: {x['documento']}" for x in socios), arquivos=arquivos,
              dados={"socios": socios, "evidencias": evidencias, "observacao": "Somente vínculos identificados na seção de sócios; não infere parentesco."})


# ── Regularize ───────────────────────────────────────────────────────────────
def _proc_regularize_primeiro(raiz, elem, ctx, arquivos=None):
    # 1) Consulta de capacidade de pagamento (PGFN)
    for arq in arquivos if arquivos is not None else candidatos(raiz, elem):
        texto = extrair_texto(arq)
        if not re.search(r"capacidade\s+de\s+pagamento", texto, re.IGNORECASE):
            continue
        m_cl = re.search(r"Classifica\S{1,3}o\s+para\s+transa\S{1,3}o\s*:?\s*([A-D]\b|omisso|sem\s+classifica[^\n]*)",
                         texto, re.IGNORECASE)
        classe = (m_cl.group(1).strip() if m_cl else "")
        m_div = re.search(r"Valor\s+total\s+da\s+d[ií]vida\s+na\s+PGFN\s+e\s+RFB\s*:\s*R\$\s*(" + _NUM + ")",
                          texto, re.IGNORECASE)
        divida = _brl(m_div.group(1)) if m_div else None
        if divida is None:
            continue
        m_cont = re.search(r"Valor\s+em\s+contencioso[^:\n]*:\s*R\$\s*(" + _NUM + ")", texto, re.IGNORECASE)
        contencioso = _brl(m_cont.group(1)) if m_cont else None

        limiar = REGULARIZE_LIMIAR
        eh = divida > limiar                      # acima de R$ 600 mil = dentro
        classe_txt = classe.upper() if classe[:1].upper() in "ABCD" and len(classe) == 1 else (classe or "sem classificação")
        valor = (f"{'SIM' if eh else 'Não'} — {'dentro' if eh else 'fora'} do Regularize | "
                 f"dívida PGFN+RFB {_fmt_brl(divida)} "
                 f"({'acima' if eh else 'não acima'} de {_fmt_brl(limiar)})")
        if contencioso:
            valor += f" | contencioso RFB {_fmt_brl(contencioso)} (não somado)"
        return _r(valor, _trecho(texto, m_div.start(), m_div.end()), [arq], status="ok",
                  dados={"regra": f"dívida PGFN+RFB > {limiar}",
                         "dentro": eh,
                         "divida_total": divida,
                         "limiar": limiar,
                         "contencioso": contencioso,
                         "classificacao_documento": classe_txt,   # informativo, não entra na regra
                         "perfil_regularize": eh})

    if arquivos is not None:
        return _r()
    # 2) Arquivo (normalmente .txt) avisando que NÃO há cadastro no Regularize
    for caminho in _varrer_pasta(raiz, ["B - Documentos", "2 - Fiscal"], set()):
        if "regularize" in os.path.basename(caminho).lower() and re.search(r"(?:sem\s+cadastro|n[aã]o\s+(?:possui|cadastrad))", extrair_texto(caminho), re.I):
            _log(f"      [✓] sem cadastro: {caminho}")
            return _r("Sem cadastro no Regularize", arquivos=[caminho], status="ok",
                      dados={"perfil_regularize": False, "sem_cadastro": True})
    return _r()


def _proc_regularize(raiz, elem, ctx):
    observacoes, candidatos_lidos = [], []
    for arq in candidatos(raiz, elem):
        candidatos_lidos.append(arq)
        texto = extrair_texto(arq)
        mj = re.search(r"CPF/CNPJ\s*:\s*(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})", texto, re.I)
        if mj and ctx.get('cnpj_matriz') and mj.group(1) != ctx['cnpj_matriz']:
            continue
        r = _proc_regularize_primeiro(raiz, elem, ctx, [arq])
        if r['valor']: observacoes.append(r)
    if observacoes:
        evidencias = [{'arquivo':r['arquivos'][0], 'trecho':r['trecho'], **r['dados']} for r in observacoes]
        if len({r['dados']['divida_total'] for r in observacoes}) > 1:
            return _r('Inconclusivo — documentos PGFN com dívidas divergentes', arquivos=candidatos_lidos,
                      status='inconclusivo', dados={'evidencias':evidencias})
        r = observacoes[0]
        r['arquivos'] = [x['arquivos'][0] for x in observacoes]
        r['dados']['evidencias'] = evidencias
        return r
    # No inferred negative from absent or unreadable capacity reports.
    if candidatos_lidos:
        return _r('Inconclusivo — dívida não extraída de documento compatível', arquivos=candidatos_lidos, status='inconclusivo')
    return _proc_regularize_primeiro(raiz, elem, ctx)


# ── Perfil previdenciário (todos os anos) ────────────────────────────────────
def _parse_perfil(texto: str) -> list[dict]:
    anos: dict[str, dict] = {}
    for linha in texto.splitlines():
        m_ano = re.search(r"\b((?:19|20)\d{2})\b", linha)
        m_st = re.search(r"\b(acima|abaixo)\b", linha, re.IGNORECASE)
        if not (m_ano and m_st):
            continue
        media = None
        m_med = re.search(r"m[eé]dia\D{0,15}(" + _NUM + ")", linha, re.IGNORECASE)
        if m_med:
            media = _brl(m_med.group(1))
        else:
            ini = m_ano.end() if m_ano.end() < m_st.start() else 0
            nums = [_brl(x) for x in re.findall(_NUM, linha[ini:m_st.start()])]
            nums = [n for n in nums if n is not None]
            media = nums[-1] if nums else None
        anos[m_ano.group(1)] = {"ano": m_ano.group(1), "media": media, "status": m_st.group(1).upper()}
    return [anos[a] for a in sorted(anos)]


def _proc_perfil(raiz, elem, ctx):
    registros, arquivos, ignorados = [], [], []
    for arq in _varrer_pasta(raiz, elem['pasta'], {'.xlsx','.pdf','.txt'}):
        encontrados = []
        if Path(arq).suffix.lower() == '.xlsx':
            import openpyxl
            try:
                with_workbook = openpyxl.load_workbook(arq,read_only=True,data_only=True)
                for sh in with_workbook:
                    cab = None
                    for numero, row in enumerate(sh.values,1):
                        nomes = [_normalizar(str(x or '')) for x in row]
                        if 'ano' in nomes and 'status' in nomes and any('media' in x for x in nomes):
                            cab = (nomes.index('ano'),next(i for i,x in enumerate(nomes) if 'media' in x),nomes.index('status'))
                            continue
                        if cab:
                            ia, im, ist = cab
                            ano = str(row[ia] or '')
                            status = _normalizar(str(row[ist] or ''))
                            media = row[im]
                            if re.fullmatch(r'(?:19|20)\d{2}',ano) and status in ('acima do limiar','abaixo do limiar','acima','abaixo') and isinstance(media,(int,float)):
                                encontrados.append({'ano':ano,'media':media,'status':'ACIMA' if status.startswith('acima') else 'ABAIXO','arquivo':arq,'planilha':sh.title,'linha':numero,'trecho':str(row)})
                with_workbook.close()
            except Exception as e:
                ignorados.append({'arquivo':arq,'motivo':str(e)})
        elif 'perfil' in _normalizar(os.path.basename(arq)):
            texto = extrair_texto(arq)
            encontrados = [{**x,'arquivo':arq,'trecho':texto[:2000]} for x in _parse_perfil(texto)]
        if encontrados:
            registros.extend(encontrados);arquivos.append(arq)
    por_ano = {}
    for x in registros: por_ano.setdefault(x['ano'],set()).add(x['status'])
    conflitos = [ano for ano,sts in por_ano.items() if len(sts)>1]
    eh = any(x['status']=='ACIMA' for x in registros) if registros and not conflitos else None
    valor = ('É PERFIL' if eh else 'Não é perfil') if eh is not None else 'Inconclusivo — sem veredito confiável ou documentos conflitantes'
    linhas = [valor] + [f"{x['ano']}: média {_fmt_num(x['media']) if x['media'] is not None else 'não identificada'} | {x['status']}" for x in registros]
    return _r('\n'.join(linhas), arquivos=arquivos, status='ok' if eh is not None else 'inconclusivo',
              dados={'eh_perfil':eh,'anos':registros,'conflitos':conflitos,'ignorados':ignorados,'evidencias':registros})


# ── Faturamento: total por ano + média ───────────────────────────────────────
_RE_TOTAL_ARQ = [
    r"TOTAL\s*DO\s*ARQUIVO[^\d]{0,20}(" + _NUM + ")",
    r"(?:receita|faturamento)\s*(?:bruta|total)?\s*[:\-]?\s*R?\$?\s*(" + _NUM + ")",
]


def _ano_do_arquivo(caminho: str, texto: str):
    anos_nome = set(re.findall(r"(?<!\d)((?:19|20)\d{2})(?!\d)", os.path.basename(caminho)))
    anos_texto = set(re.findall(r"(?:ano[\s-]*calend[aá]rio|exerc[ií]cio|\bano)\D{0,15}((?:19|20)\d{2})", texto, re.I))
    anos = anos_nome | anos_texto
    return next(iter(anos)) if len(anos) == 1 else None


def _proc_faturamento(raiz, elem, ctx):
    arquivos = encontrar_arquivos(raiz, elem)
    if not arquivos:
        return _r()
    por_ano: dict[str, float] = {}
    evidencias = []
    conflitos = set()
    def registrar(ano, valor, arq, trecho):
        evidencias.append({"ano": ano, "valor": valor, "arquivo": arq, "trecho": trecho})
        if ano in por_ano and por_ano[ano] != valor:
            conflitos.add(ano)
        por_ano[ano] = valor
    usados: list[str] = []
    for arq in arquivos:
        if Path(arq).suffix.lower() == '.xlsx':
            import openpyxl
            try:
                wb = openpyxl.load_workbook(arq,read_only=True,data_only=True)
                for sh in wb:
                    cab = None
                    for numero,row in enumerate(sh.values,1):
                        nomes = [_normalizar(str(x or '')) for x in row]
                        if 'ano' in nomes:
                            cols = [i for i,x in enumerate(nomes) if x in ('faturamento','faturamento anual','receita bruta','receita anual','faturamento total')]
                            cab = (nomes.index('ano'),cols[0]) if len(cols)==1 else None
                            continue
                        if cab:
                            ano = str(row[cab[0]] or '')
                            v = row[cab[1]]
                            valor = float(v) if isinstance(v,(int,float)) else _brl(str(v))
                            if re.fullmatch(r'(?:19|20)\d{2}',ano) and valor is not None and valor >= 0:
                                registrar(ano,valor,arq,f'{sh.title}!linha {numero}: {row}')
                                if arq not in usados: usados.append(arq)
                wb.close()
            except Exception as e:
                _log(f'[faturamento] planilha não lida: {e}')
            continue
        texto = extrair_texto(arq)
        if re.search(r"DARF|Per[ií]odo Apura[çc][aã]o|SOMA VALOR ORIGINAL|M[ÉE]DIA.*DOCS", texto, re.I):
            continue  # arrecadação e DCTF não representam receita anual
        # Um arquivo com várias linhas "ANO ... VALOR"
        linhas_ano = re.findall(r"(?m)^\s*((?:19|20)\d{2})\D{1,40}?(" + _NUM + r")\s*$", texto)
        if linhas_ano and re.search(r"faturamento|receita", texto, re.I):
            for ano, v in linhas_ano:
                if _brl(v) is not None:
                    registrar(ano, _brl(v), arq, f"{ano}: {v}")
            usados.append(arq)
            continue

        if re.search(r"(?:faturamento|receita)\s+mensal|(?:janeiro|fevereiro|marco|março|abril|maio|junho|julho|agosto|setembro|outubro|novembro|dezembro)\s+(?:de\s+)?20\d{2}", texto, re.I):
            continue
        total = None
        for p in _RE_TOTAL_ARQ:
            m = re.search(p, texto, re.IGNORECASE)
            if m and _brl(m.group(1)) is not None:
                total = _brl(m.group(1))
                break
        if total is None:
            continue
        ano = _ano_do_arquivo(arq, texto)
        if not ano:
            continue
        registrar(ano, total, arq, m.group(0))
        usados.append(arq)

    if not por_ano:
        return _r(arquivos=arquivos, status="sem_dado", dados={"motivo":"Não foi possível identificar totais anuais confiáveis; não é faturamento zero."})

    ano_atual = datetime.now().year
    fechados = {a:v for a,v in por_ano.items() if str(a).isdigit() and int(a) < ano_atual and a not in conflitos}
    media = sum(fechados.values()) / len(fechados) if fechados and not any(str(a).isdigit() and int(a) < ano_atual for a in conflitos) else None
    linhas = [f"{a}: {_fmt_brl(v)}" for a, v in sorted(por_ano.items())]
    linhas.append("Média dos anos fechados: " + (_fmt_brl(media) if media is not None else "Inconclusivo (sem anos fechados ou valores conflitantes)"))
    linhas.append("Anos usados na média: " + (", ".join(sorted(fechados)) if media is not None else "nenhum"))
    if conflitos:
        linhas.append("Conflitos por ano: " + ", ".join(sorted(conflitos)))
    return _r("\n".join(linhas), "", usados,
              dados={"por_ano": por_ano, "media": media, "anos_media": sorted(fechados) if media is not None else [], "conflitos": sorted(conflitos), "evidencias": evidencias})


# ── Derivados do faturamento: Fastpass / Supply / Finap ──────────────────────
# Lista explícita extraída da planilha cnae Finep (2).xlsx. Normalização: 7 dígitos.
CNAES_FINAP = {'0141501': 'Produção de sementes certificadas, exceto de forrageiras para pasto',
 '0141502': 'Produção de sementes certificadas de forrageiras para formação de pasto',
 '0142300': 'Produção de mudas e outras formas de propagação vegetal, certificadas',
 '0155502': 'Produção de pintos de um dia',
 '0155505': 'Produção de ovos',
 '0891600': 'Extração de minerais para fabricação de adubos, fertilizantes e outros produtos químicos',
 '1011201': 'Frigorífico – abate de bovinos',
 '1011202': 'Frigorífico – abate de equinos',
 '1011203': 'Frigorífico – abate de ovinos e caprinos',
 '1011204': 'Frigorífico – abate de bufalinos',
 '1012103': 'Frigorífico – abate de suínos',
 '1012104': 'Matadouro – abate de suínos sob contrato',
 '1013901': 'Fabricação de produtos de carne',
 '1013902': 'Preparação de subprodutos do abate',
 '1020102': 'Fabricação de conservas de peixes, crustáceos e moluscos',
 '1031700': 'Fabricação de conservas de frutas',
 '1032501': 'Fabricação de conservas de palmito',
 '1032599': 'Fabricação de conservas de legumes e outros vegetais, exceto palmito',
 '1033301': 'Fabricação de sucos concentrados de frutas, hortaliças e legumes',
 '1033302': 'Fabricação de sucos de frutas, hortaliças e legumes, exceto concentrados',
 '1041400': 'Fabricação de óleos vegetais em bruto, exceto óleo de milho',
 '1042200': 'Fabricação de óleos vegetais refinados, exceto óleo de milho',
 '1043100': 'Fabricação de margarina e outras gorduras vegetais e de óleos não comestíveis de animais',
 '1052000': 'Fabricação de laticínios',
 '1053800': 'Fabricação de sorvetes e outros gelados comestíveis',
 '1061902': 'Fabricação de produtos do arroz',
 '1063500': 'Fabricação de farinha de mandioca e derivados',
 '1064300': 'Fabricação de farinha de milho e derivados, exceto óleos de milho',
 '1065101': 'Fabricação de amidos e féculas de vegetais',
 '1065102': 'Fabricação de óleo de milho em bruto',
 '1065103': 'Fabricação de óleo de milho refinado',
 '1066000': 'Fabricação de alimentos para animais – ração',
 '1071600': 'Fabricação de açúcar em bruto',
 '1072401': 'Fabricação de açúcar de cana refinado',
 '1072402': 'Fabricação de açúcar de cereais e de beterraba',
 '1082100': 'Fabricação de produtos à base de café',
 '1091101': 'Fabricação de produtos de panificação industrial',
 '1091102': 'Fabricação de produtos de padaria e confeitaria com predominância de produção própria',
 '1092900': 'Fabricação de biscoitos e bolachas',
 '1093701': 'Fabricação de produtos derivados do cacau e de chocolates',
 '1093702': 'Fabricação de frutas cristalizadas, balas e semelhantes',
 '1094500': 'Fabricação de massas alimentícias',
 '1095300': 'Fabricação de especiarias, molhos, temperos e condimentos',
 '1096100': 'Fabricação de alimentos e pratos prontos',
 '1099602': 'Fabricação de pós alimentícios',
 '1099603': 'Fabricação de fermentos e leveduras',
 '1099605': 'Fabricação de produtos para infusão',
 '1099606': 'Fabricação de adoçantes naturais e artificiais',
 '1099607': 'Fabricação de alimentos dietéticos e complementos alimentares',
 '1099699': 'Fabricação de outros produtos alimentícios não especificados anteriormente',
 '1710900': 'Fabricação de celulose e outras pastas para a fabricação de papel',
 '1721400': 'Fabricação de papel',
 '1722200': 'Fabricação de cartolina e papel-cartão',
 '1731100': 'Fabricação de embalagens de papel',
 '1732000': 'Fabricação de embalagens de cartolina e papel-cartão',
 '1733800': 'Fabricação de chapas e embalagens de papelão ondulado',
 '1931400': 'Fabricação de álcool',
 '1932200': 'Fabricação de biocombustíveis, exceto álcool',
 '2011800': 'Fabricação de cloro e álcalis',
 '2012600': 'Fabricação de intermediários para fertilizantes',
 '2013401': 'Fabricação de adubos e fertilizantes organominerais',
 '2013402': 'Fabricação de adubos e fertilizantes, exceto organominerais',
 '2014200': 'Fabricação de gases industriais',
 '2019301': 'Elaboração de combustíveis nucleares',
 '2019399': 'Fabricação de outros produtos químicos inorgânicos não especificados anteriormente',
 '2021500': 'Fabricação de produtos petroquímicos básicos',
 '2022300': 'Fabricação de intermediários para plastificantes, resinas e fibras',
 '2029100': 'Fabricação de produtos químicos orgânicos não especificados anteriormente',
 '2031200': 'Fabricação de resinas termoplásticas',
 '2032100': 'Fabricação de resinas termofixas',
 '2033900': 'Fabricação de elastômeros',
 '2040100': 'Fabricação de fibras artificiais e sintéticas',
 '2051700': 'Fabricação de defensivos agrícolas',
 '2052500': 'Fabricação de desinfestantes domissanitários',
 '2061400': 'Fabricação de sabões e detergentes sintéticos',
 '2062200': 'Fabricação de produtos de limpeza e polimento',
 '2063100': 'Fabricação de cosméticos, produtos de perfumaria e de higiene pessoal',
 '2071100': 'Fabricação de tintas, vernizes, esmaltes e lacas',
 '2072000': 'Fabricação de tintas de impressão',
 '2073800': 'Fabricação de impermeabilizantes, solventes e produtos afins',
 '2091600': 'Fabricação de adesivos e selantes',
 '2092401': 'Fabricação de pólvoras, explosivos e detonantes',
 '2092402': 'Fabricação de artigos pirotécnicos',
 '2092403': 'Fabricação de fósforos de segurança',
 '2093200': 'Fabricação de aditivos de uso industrial',
 '2094100': 'Fabricação de catalisadores',
 '2099101': 'Fabricação de chapas, filmes, papéis e outros materiais e produtos químicos para fotografia',
 '2099199': 'Fabricação de outros produtos químicos não especificados anteriormente',
 '2110600': 'Fabricação de produtos farmoquímicos',
 '2121101': 'Fabricação de medicamentos alopáticos para uso humano',
 '2121102': 'Fabricação de medicamentos homeopáticos para uso humano',
 '2121103': 'Fabricação de medicamentos fitoterápicos para uso humano',
 '2122000': 'Fabricação de medicamentos para uso veterinário',
 '2123800': 'Fabricação de preparações farmacêuticas',
 '2229302': 'Fabricação de artefatos de material plástico para usos industriais',
 '2229303': 'Fabricação de artefatos de material plástico para uso na construção',
 '2229399': 'Fabricação de artefatos de material plástico para outros usos',
 '2311700': 'Fabricação de vidro plano e de segurança',
 '2319200': 'Fabricação de artigos de vidro',
 '2320600': 'Fabricação de cimento',
 '2349499': 'Fabricação de produtos cerâmicos não refratários não especificados anteriormente',
 '2421100': 'Produção de semiacabados de aço',
 '2422901': 'Produção de laminados planos de aço ao carbono, revestidos ou não',
 '2422902': 'Produção de laminados planos de aços especiais',
 '2423701': 'Produção de tubos de aço sem costura',
 '2423702': 'Produção de laminados longos de aço, exceto tubos',
 '2424501': 'Produção de arames de aço',
 '2424502': 'Produção de relaminados, trefilados e perfilados de aço, exceto arames',
 '2431800': 'Produção de tubos de aço com costura',
 '2439300': 'Produção de outros tubos de ferro e aço',
 '2441501': 'Produção de alumínio e suas ligas em formas primárias',
 '2441502': 'Produção de laminados de alumínio',
 '2449101': 'Produção de zinco em formas primárias',
 '2449102': 'Produção de laminados de zinco',
 '2511000': 'Fabricação de estruturas metálicas',
 '2513600': 'Fabricação de obras de caldeiraria pesada',
 '2521700': 'Fabricação de tanques, reservatórios metálicos e caldeiras para aquecimento central',
 '2522500': 'Fabricação de caldeiras geradoras de vapor',
 '2531401': 'Produção de forjados de aço',
 '2531402': 'Produção de forjados de metais não ferrosos e suas ligas',
 '2532201': 'Produção de artefatos estampados de metal',
 '2543800': 'Fabricação de ferramentas',
 '2599399': 'Fabricação de outros produtos de metal não especificados anteriormente',
 '2610800': 'Fabricação de componentes eletrônicos',
 '2621300': 'Fabricação de equipamentos de informática',
 '2622100': 'Fabricação de periféricos para equipamentos de informática',
 '2631100': 'Fabricação de equipamentos transmissores de comunicação, peças e acessórios',
 '2632900': 'Fabricação de aparelhos telefônicos e de outros equipamentos de comunicação',
 '2651500': 'Fabricação de aparelhos e equipamentos de medida, teste e controle',
 '2660400': 'Fabricação de aparelhos eletromédicos e eletroterapêuticos e equipamentos de irradiação',
 '2670101': 'Fabricação de equipamentos e instrumentos ópticos, peças e acessórios',
 '2670102': 'Fabricação de aparelhos fotográficos e cinematográficos, peças e acessórios',
 '2710401': 'Fabricação de geradores de corrente contínua e alternada, peças e acessórios',
 '2710402': 'Fabricação de transformadores, indutores, conversores, sincronizadores e semelhantes, peças e '
            'acessórios',
 '2710403': 'Fabricação de motores elétricos, peças e acessórios',
 '2721000': 'Fabricação de pilhas, baterias e acumuladores elétricos, exceto para veículos automotores',
 '2722801': 'Fabricação de baterias e acumuladores para veículos automotores',
 '2731700': 'Fabricação de aparelhos e equipamentos para distribuição e controle de energia elétrica',
 '2732500': 'Fabricação de material elétrico para instalações em circuito de consumo',
 '2733300': 'Fabricação de fios, cabos e condutores elétricos isolados',
 '2790202': 'Fabricação de equipamentos para sinalização e alarme',
 '2790299': 'Fabricação de outros equipamentos e aparelhos elétricos não especificados anteriormente',
 '2811900': 'Fabricação de motores e turbinas, peças e acessórios',
 '2812700': 'Fabricação de equipamentos hidráulicos e pneumáticos, peças e acessórios',
 '2813500': 'Fabricação de válvulas, registros e dispositivos semelhantes, peças e acessórios',
 '2814301': 'Fabricação de compressores para uso industrial, peças e acessórios',
 '2814302': 'Fabricação de compressores para uso não industrial, peças e acessórios',
 '2815101': 'Fabricação de rolamentos para fins industriais',
 '2815102': 'Fabricação de equipamentos de transmissão para fins industriais, exceto rolamentos',
 '2821601': 'Fabricação de fornos industriais e equipamentos não elétricos para instalações térmicas',
 '2821602': 'Fabricação de estufas e fornos elétricos para fins industriais',
 '2822401': 'Fabricação de máquinas e equipamentos para transporte e elevação de pessoas',
 '2822402': 'Fabricação de máquinas e equipamentos para transporte e elevação de cargas',
 '2823200': 'Fabricação de máquinas e aparelhos de refrigeração e ventilação para uso industrial e comercial',
 '2824101': 'Fabricação de aparelhos e equipamentos de ar condicionado para uso industrial',
 '2824102': 'Fabricação de aparelhos e equipamentos de ar condicionado para uso não industrial',
 '2825900': 'Fabricação de máquinas e equipamentos para saneamento básico e ambiental',
 '2829101': 'Fabricação de máquinas de escrever, calcular e outros equipamentos não eletrônicos para '
            'escritório, peças e acessórios',
 '2829199': 'Fabricação de outras máquinas e equipamentos de uso geral não especificados anteriormente',
 '2831300': 'Fabricação de tratores agrícolas, peças e acessórios',
 '2832100': 'Fabricação de equipamentos para irrigação agrícola',
 '2833000': 'Fabricação de máquinas e equipamentos para agricultura e pecuária, peças e acessórios',
 '2840200': 'Fabricação de máquinas-ferramenta, peças e acessórios',
 '2851800': 'Fabricação de máquinas e equipamentos para prospecção e extração de petróleo',
 '2852600': 'Fabricação de máquinas e equipamentos para extração mineral',
 '2853400': 'Fabricação de tratores, peças e acessórios, exceto agrícolas',
 '2854200': 'Fabricação de máquinas e equipamentos para terraplenagem, pavimentação e construção',
 '2861500': 'Fabricação de máquinas para a indústria metalúrgica, peças e acessórios, exceto '
            'máquinas-ferramenta',
 '2862300': 'Fabricação de máquinas e equipamentos para as indústrias de alimentos, bebidas e fumo, peças e '
            'acessórios',
 '2863100': 'Fabricação de máquinas e equipamentos para a indústria têxtil, peças e acessórios',
 '2864000': 'Fabricação de máquinas e equipamentos para as indústrias do vestuário, couro e calçados, peças '
            'e acessórios',
 '2865800': 'Fabricação de máquinas e equipamentos para as indústrias de celulose, papel e papelão e '
            'artefatos, peças e acessórios',
 '2866600': 'Fabricação de máquinas e equipamentos para a indústria do plástico, peças e acessórios',
 '2869100': 'Fabricação de máquinas e equipamentos para uso industrial específico',
 '2910702': 'Fabricação de chassis com motor para automóveis, camionetas e utilitários',
 '2910703': 'Fabricação de motores para automóveis, camionetas e utilitários',
 '2920402': 'Fabricação de motores para caminhões e ônibus',
 '2930101': 'Fabricação de cabines, carrocerias e reboques para caminhões',
 '2930102': 'Fabricação de carrocerias para ônibus',
 '2930103': 'Fabricação de cabines, carrocerias e reboques para outros veículos automotores',
 '2941700': 'Fabricação de peças e acessórios para o sistema motor de veículos automotores',
 '2942500': 'Fabricação de peças e acessórios para os sistemas de marcha e transmissão de veículos '
            'automotores',
 '2943300': 'Fabricação de peças e acessórios para o sistema de freios de veículos automotores',
 '2944100': 'Fabricação de peças e acessórios para o sistema de direção e suspensão de veículos automotores',
 '2945000': 'Fabricação de material elétrico e eletrônico para veículos automotores, exceto baterias',
 '2949299': 'Fabricação de outras peças e acessórios para veículos automotores não especificadas '
            'anteriormente',
 '3011301': 'Construção de embarcações de grande porte',
 '3011302': 'Construção de embarcações para uso comercial e usos especiais',
 '3012100': 'Construção de embarcações para esporte e lazer',
 '3031800': 'Fabricação de locomotivas, vagões e outros materiais rodantes',
 '3032600': 'Fabricação de peças e acessórios para veículos ferroviários',
 '3041500': 'Fabricação de aeronaves',
 '3042300': 'Fabricação de turbinas, motores e outros componentes e peças para aeronaves',
 '3250701': 'Fabricação de instrumentos não eletrônicos e utensílios para uso médico, cirúrgico, '
            'odontológico e de laboratório',
 '3250702': 'Fabricação de mobiliário para uso médico, cirúrgico, odontológico e de laboratório',
 '3250703': 'Fabricação de aparelhos e utensílios para correção de defeitos físicos e aparelhos ortopédicos '
            'em geral sob encomenda',
 '3250704': 'Fabricação de aparelhos e utensílios para correção de defeitos físicos e aparelhos ortopédicos '
            'em geral, exceto sob encomenda',
 '3250705': 'Fabricação de materiais para medicina e odontologia',
 '3250707': 'Fabricação de artigos ópticos',
 '3250708': 'Fabricação de artefatos de tecido não tecido para uso odonto-médico-hospitalar',
 '3520401': 'Produção de gás; processamento de gás natural',
 '3530100': 'Produção e distribuição de vapor, água quente e ar condicionado',
 '3831901': 'Recuperação de sucatas de alumínio',
 '3831999': 'Recuperação de materiais metálicos, exceto alumínio',
 '3832700': 'Recuperação de materiais plásticos',
 '3839401': 'Usinas de compostagem',
 '3839499': 'Recuperação de materiais não especificados anteriormente',
 '6201501': 'Desenvolvimento de programas de computador sob encomenda',
 '6202300': 'Desenvolvimento e licenciamento de programas de computador customizáveis',
 '6203100': 'Desenvolvimento e licenciamento de programas de computador não customizáveis',
 '7120100': 'Testes e análises técnicas',
 '7210000': 'Pesquisa e desenvolvimento experimental em ciências físicas e naturais',
 '7220700': 'Pesquisa e desenvolvimento experimental em ciências sociais e humanas',
 '8640202': 'Laboratórios clínicos e biologia molecular'}

def _eh_industria(cnae: str) -> bool:
    return re.sub(r"\D", "", cnae or "") in CNAES_FINAP


def _calcular_derivado(elem: dict, resultados: list[dict], ctx: dict) -> str:
    fat = next((r for r in resultados if r["id"] == "faturamento"), None)
    media = (fat or {}).get("dados", {}).get("media")
    if media is None:
        return "Inconclusivo — sem faturamento de anos fechados"
    limiar = elem["limiar"]
    if elem["id"] == "fastpass":
        return f"{'Sim' if media > limiar else 'Inconclusivo'} — média {_fmt_brl(media)}; exige acima de {_fmt_brl(limiar)}"
    if media < limiar or (elem["id"] == "finap" and media == limiar):
        return f"Não — média {_fmt_brl(media)} (abaixo de {_fmt_brl(limiar)})"
    if elem.get("exige_industria"):
        cnae = ctx.get("cnae_principal")
        if not cnae:
            return f"Verificar — média {_fmt_brl(media)}, mas o CNAE principal não foi identificado (precisa constar na lista Finap)"
        if not _eh_industria(cnae):
            return f"Não — média {_fmt_brl(media)}, mas o CNAE principal ({cnae}) não está na lista de CNAEs Finap"
        return f"Sim — média {_fmt_brl(media)} | CNAE permitido {cnae}"
    return f"Sim — média {_fmt_brl(media)} (a partir de {_fmt_brl(limiar)})"


PROCESSADORES = {
    "existencia":    _proc_existencia,
    "manual":        _proc_manual,
    "honorarios":    _proc_honorarios,
    "gross_up":      _proc_gross_up,
    "booleano":      _proc_booleano,
    "data_abertura": _proc_data_abertura,
    "simples":       _proc_simples,
    "qsa":           _proc_qsa,
    "regularize":    _proc_regularize,
    "perfil":        _proc_perfil,
    "faturamento":   _proc_faturamento,
}


# ══════════════════════════════════════════════════════════════════════════════
#  ANÁLISE PRINCIPAL
# ══════════════════════════════════════════════════════════════════════════════
def analisar_diretorio(raiz: str, callback=None) -> list[dict]:
    _log(f"\n{'═'*60}")
    _log("  ANÁLISE INICIADA")
    _log(f"  Pasta: {raiz}")
    _log(f"{'═'*60}")
    resultados: list[dict] = []
    ctx: dict = {}
    total = len(ELEMENTOS)

    for i, elemento in enumerate(ELEMENTOS):
        _log(f"\n  [{i+1:02d}/{total}] {elemento['campo']}")
        if callback:
            callback(i, total, elemento["campo"])

        base = {
            "id":        elemento["id"],
            "categoria": elemento["categoria"],
            "campo":     elemento["campo"],
            "descricao": elemento["descricao"],
            "tipo":      elemento.get("tipo", "simples"),
        }

        tipo = elemento.get("tipo")
        if tipo == "derivado":
            valor = _calcular_derivado(elemento, resultados, ctx)
            fat = next((r for r in resultados if r["id"] == "faturamento"), {})
            parcial = _r(valor, arquivos=fat.get("arquivos", []), dados={"faturamento": fat.get("dados", {}), "limiar": elemento["limiar"], "cnae_principal": ctx.get("cnae_principal"), "cnae_arquivo": ctx.get("cnae_arquivo"), "cnae_trecho": ctx.get("cnae_trecho"), "cnae_na_lista": _eh_industria(ctx.get("cnae_principal", ""))})
        else:
            try:
                parcial = PROCESSADORES[tipo](raiz, elemento, ctx)
            except Exception as e:                      # um elemento nunca derruba os demais
                _log(f"      [ERRO] {elemento['id']}: {e}")
                parcial = _r(status="erro", trecho=str(e))
        if parcial["valor"]:
            _log(f"      → {parcial['valor'].splitlines()[0][:90]!r}")
        resultados.append({**base, **parcial})

    ok = sum(1 for r in resultados if r["status"] == "ok")
    _log(f"\n{'═'*60}")
    _log(f"  ANÁLISE CONCLUÍDA — {ok}/{total} encontrados")
    _log(f"{'═'*60}\n")
    if callback:
        callback(total, total, "Concluído")
    for r in resultados:
        if r["id"] in {'socios': 'Perfil Family', 'regularize': 'Perfil Transação', 'perfil': 'Perfil Previdenciário', 'fastpass': 'Perfil Fastpass', 'supply': 'Perfil Supply', 'finap': 'Perfil Finap - Lei do Bem'}:
            r["campo"] = {'socios': 'Perfil Family', 'regularize': 'Perfil Transação', 'perfil': 'Perfil Previdenciário', 'fastpass': 'Perfil Fastpass', 'supply': 'Perfil Supply', 'finap': 'Perfil Finap - Lei do Bem'}[r["id"]]
    return resultados


# ══════════════════════════════════════════════════════════════════════════════
#  CHECKLIST DE ALERTAS
#  (Procuração não gera alerta: vem do Project)
# ══════════════════════════════════════════════════════════════════════════════
def _build_checklist(resultados: list[dict]) -> list[dict]:
    def _val(id_elem: str) -> str:
        r = next((x for x in resultados if x["id"] == id_elem), None)
        return (r.get("valor") or "") if r else ""

    fp        = _val("fastpass").lower()
    baixo_fat = False  # Fastpass abaixo do limite é inconclusivo, não reprovação.

    itens = [
        (not _val("contrato_existe"),               "SEM CONTRATO"),
        (baixo_fat,                                  "FATURAMENTO MÉDIO ABAIXO DE R$ 200M"),
        (_val("regularize").upper().startswith("SIM"), "COM REGULARIZE"),
    ]
    return [{"alerta": bool(a), "descricao": d} for a, d in itens]

# ══════════════════════════════════════════════════════════════════════════════
#  LOG (stderr, para não poluir o JSON no stdout)
# ══════════════════════════════════════════════════════════════════════════════
def _log(msg: str) -> None:
    print(msg, file=sys.stderr, flush=True)


# ══════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════
def main() -> None:
    if len(sys.argv) < 2:
        print(json.dumps({
            "erro": "Uso: python3 analisar_job.py <pasta_raiz>"
        }), flush=True)
        sys.exit(1)

    pasta_raiz = sys.argv[1]

    if not os.path.isdir(pasta_raiz):
        print(json.dumps({
            "erro": f"Diretório não encontrado: {pasta_raiz}"
        }), flush=True)
        sys.exit(1)

    def callback(atual, total, campo):
        _log(f"[{atual}/{total}] {campo}")

    try:
        resultados = analisar_diretorio(pasta_raiz, callback)
        checklist  = _build_checklist(resultados)

        # Resumo numérico
        ok       = sum(1 for r in resultados if r["status"] == "ok")
        alertas  = sum(1 for r in resultados if r["status"] in ("nao_encontrado", "sem_dado") and not r.get("valor"))
        n_alertas = sum(1 for c in checklist if c["alerta"])

        saida = {
            "pasta_raiz":  pasta_raiz,
            "analisado_em": datetime.now().isoformat(),
            "resumo": {
                "total":         len(resultados),
                "encontrados":   ok,
                "nao_encontrados": alertas,
                "alertas_checklist": n_alertas,
            },
            "checklist":   checklist,
            "resultados":  resultados,
        }

        print(json.dumps(saida, ensure_ascii=False, default=str), flush=True)

    except Exception as e:
        import traceback
        print(json.dumps({
            "erro":      str(e),
            "traceback": traceback.format_exc(),
        }), flush=True)
        sys.exit(1)


if __name__ == "__main__":
    main()
